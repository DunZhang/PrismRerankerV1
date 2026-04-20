"""Summarize the enriched JSONL produced by ``evaluate.py`` into a single xlsx.

Each row in the output corresponds to one input sample; columns are the metric
names; the final row holds the per-column average over non-empty cells.

Usage:
    uv run python -m evaluate_relevance_contribution_evidence.summarize \\
        --input  /mnt/g/PrismRerankerV1Data/.../result_eval.jsonl
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from evaluate_relevance_contribution_evidence.evaluate import SCORE_FIELDS

log = logging.getLogger("evaluate_relevance_contribution_evidence_summarize")

DEFAULT_INPUT_PATH = Path(
    "/mnt/g/PrismRerankerV1Data/relevance_contribution_evidence_evaluate_result/"
    "label_eval.jsonl"
)

METRIC_COLUMNS: tuple[str, ...] = (
    "label_match",
    "format_score",
    "entity_fidelity",
    *SCORE_FIELDS,
)
SHEET_TITLE = "metrics"
HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center")
COLUMN_WIDTH = 22
FLOAT_FORMAT = "0.0000"


def _default_output(input_path: Path) -> Path:
    return input_path.parent / f"{input_path.stem}_summary.xlsx"


def _setup_logging() -> None:
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)-5s | %(message)s",
        datefmt="%H:%M:%S",
        stream=sys.stderr,
    )


def _load_rows(input_path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with open(input_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            rows.append(json.loads(line))
    return rows


def _row_to_metrics(row: dict[str, Any]) -> dict[str, float | int | None]:
    """Map one JSONL row to {metric_name: value-or-None}."""
    metrics: dict[str, float | int | None] = {col: None for col in METRIC_COLUMNS}

    metrics["label_match"] = 1 if row.get("label_match") == "yes" else 0

    fmt = row.get("format_score")
    if isinstance(fmt, (int, float)):
        metrics["format_score"] = float(fmt)

    ef = row.get("entity_fidelity")
    if isinstance(ef, dict) and isinstance(ef.get("score"), (int, float)):
        metrics["entity_fidelity"] = float(ef["score"])

    if row.get("eval_status") == "scored":
        scores = row.get("eval_scores")
        if isinstance(scores, dict):
            for field in SCORE_FIELDS:
                v = scores.get(field)
                if isinstance(v, (int, float)):
                    metrics[field] = int(v)
    return metrics


def _column_average(
    rows_metrics: list[dict[str, float | int | None]], column: str
) -> float | None:
    vals = [m[column] for m in rows_metrics if m[column] is not None]
    if not vals:
        return None
    return sum(vals) / len(vals)


def _write_xlsx(
    rows_metrics: list[dict[str, float | int | None]], output_path: Path
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    for col_idx, name in enumerate(METRIC_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = CENTER

    for row_idx, metrics in enumerate(rows_metrics, start=2):
        for col_idx, name in enumerate(METRIC_COLUMNS, start=1):
            value = metrics[name]
            if value is None:
                continue
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, float):
                cell.number_format = FLOAT_FORMAT
            cell.alignment = CENTER

    avg_row = len(rows_metrics) + 2
    for col_idx, name in enumerate(METRIC_COLUMNS, start=1):
        avg = _column_average(rows_metrics, name)
        if avg is None:
            continue
        cell = ws.cell(row=avg_row, column=col_idx, value=float(avg))
        cell.number_format = FLOAT_FORMAT
        cell.font = BOLD
        cell.alignment = CENTER

    for col_idx in range(1, len(METRIC_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTH

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize pred-quality eval JSONL.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT_PATH)
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output xlsx path. Defaults to <input_parent>/<input_stem>_summary.xlsx.",
    )
    args = parser.parse_args()

    _setup_logging()

    if not args.input.exists():
        log.error("input not found: %s", args.input)
        sys.exit(1)

    output_path: Path = args.output or _default_output(args.input)

    rows = _load_rows(args.input)
    log.info("Loaded %d rows from %s", len(rows), args.input)

    rows_metrics = [_row_to_metrics(r) for r in rows]
    _write_xlsx(rows_metrics, output_path)
    log.info("Wrote %s", output_path)

    log.info("-" * 60)
    log.info("Per-column averages (over non-empty cells):")
    for name in METRIC_COLUMNS:
        avg = _column_average(rows_metrics, name)
        n = sum(1 for m in rows_metrics if m[name] is not None)
        if avg is None:
            log.info("  %-26s  n=%-5d  (no values)", name, n)
        else:
            log.info("  %-26s  n=%-5d  avg=%.4f", name, n, avg)


if __name__ == "__main__":
    main()
